import asyncio
from io import BytesIO

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas
from starlette.datastructures import UploadFile

from app.config import Settings
from app.errors import AppError
from app.parsing import Section, chunk_sections, parse
from app.providers import CompatibleProvider, MockProvider, tokenize
from app.schemas import Answer
from app.storage import LocalStorage


@pytest.mark.parametrize("suffix", [".md", ".txt", ".docx", ".pdf"])
def test_four_formats(tmp_path, suffix):
    path = tmp_path / ("guide" + suffix)
    if suffix == ".docx":
        doc = Document()
        doc.add_heading("校园网络", 1)
        doc.add_paragraph("连接 Campus-WiFi，错误代码 809。")
        doc.save(path)
    elif suffix == ".pdf":
        canvas = Canvas(str(path))
        canvas.drawString(50, 750, "VPN error 809")
        canvas.showPage()
        canvas.drawString(50, 750, "Check network")
        canvas.save()
    else:
        path.write_text("# 校园网络\n\n连接 Campus-WiFi，错误代码 809。", encoding="utf8")
    sections = parse(path, path.name)
    assert any("809" in s.content for s in sections)
    if suffix == ".pdf":
        assert [s.page for s in sections] == [1, 2]
    if suffix in {".md", ".docx"}:
        assert any("校园网络" in s.heading for s in sections)


@pytest.mark.parametrize(
    "name,content,code", [("empty.txt", b"", "EMPTY_DOCUMENT"), ("broken.docx", b"not a zip", "PARSE_FAILED")]
)
def test_invalid_files(tmp_path, name, content, code):
    p = tmp_path / name
    p.write_bytes(content)
    with pytest.raises(AppError) as exc:
        parse(p, name)
    assert exc.value.code == code


def test_scan_and_pdf_page_limit(tmp_path):
    p = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(400, 600)
    writer.write(p)
    with pytest.raises(AppError, match="首版") as exc:
        parse(p, p.name)
    assert exc.value.code == "OCR_REQUIRED"
    for _ in range(100):
        writer.add_blank_page(400, 600)
    writer.write(p)
    with pytest.raises(AppError) as exc:
        parse(p, p.name)
    assert exc.value.code == "PDF_TOO_LONG"


def test_chunk_boundaries_metadata_and_overlap():
    content = "网络操作步骤。" * 350
    chunks = chunk_sections([Section(content, "一级 / 二级", page=3, paragraph=2)])
    assert all(len(c.content) <= 1200 for c in chunks)
    assert all(c.heading == "一级 / 二级" and c.page == 3 for c in chunks)
    assert chunks[0].content[-120:] == chunks[1].content[:120]
    assert len(chunks) < 10


@pytest.mark.asyncio
async def test_storage_limit_and_traversal(tmp_path, monkeypatch):
    import app.storage as module

    monkeypatch.setattr(module, "settings", lambda: Settings(storage_dir=tmp_path, max_upload_bytes=4))
    store = LocalStorage()
    with pytest.raises(AppError):
        store.path("../secret")
    with pytest.raises(AppError) as exc:
        await store.save(UploadFile(BytesIO(b"12345"), filename="guide.txt"))
    assert exc.value.code == "FILE_TOO_LARGE"
    assert list(tmp_path.iterdir()) == []


@pytest.mark.asyncio
async def test_mock_determinism_and_refusal():
    a, b = MockProvider(), MockProvider()
    assert await a.embed(["校园 VPN 809"]) == await b.embed(["校园 VPN 809"])
    assert "809" in tokenize("VPN 错误代码809") and "vpn" in tokenize("VPN 错误代码809")
    source = dict(number=1, content="VPN 错误809，请检查网络。", heading="VPN", document_name="指南")
    assert (await a.answer("VPN 809 如何处理", [source], [])).sources == [1]
    assert not (await a.answer("火星旅游", [source], [])).answerable


def test_real_configuration_and_index_identity(monkeypatch):
    for name in Settings.model_fields:
        monkeypatch.delenv(name.upper(), raising=False)
    with pytest.raises(ValueError) as secret_error:
        Settings(model_mode="real", llm_api_key="secret-must-not-appear")
    assert "secret-must-not-appear" not in str(secret_error.value)
    with pytest.raises(ValueError):
        Settings(model_mode="real")
    with pytest.raises(ValueError):
        Settings(embedding_dimensions=768)
    cfg = Settings(
        model_mode="real",
        llm_api_base="http://invalid",
        llm_api_key="test",
        llm_model="chat",
        embedding_api_base="http://invalid",
        embedding_api_key="test",
        embedding_model="embed",
    )
    assert cfg.index_key != Settings().index_key


@pytest.mark.asyncio
async def test_embedding_dimension_and_invalid_structured_output(monkeypatch):
    model = CompatibleProvider()

    async def embedding(*args):
        return {"data": [{"index": 0, "embedding": [1.0] * 3}]}

    monkeypatch.setattr(model, "request", embedding)
    with pytest.raises(AppError) as exc:
        await model.embed(["VPN"])
    assert exc.value.code == "EMBEDDING_OUTPUT"

    async def chat(*args):
        return {"choices": [{"message": {"content": "not json"}}]}

    monkeypatch.setattr(model, "request", chat)
    with pytest.raises(AppError) as exc:
        await model.structured("system", {}, Answer)
    assert exc.value.code == "MODEL_OUTPUT"


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", ["timeout", "rate_limit", "unauthorized"])
async def test_network_failures_have_bounded_retries(monkeypatch, failure):
    calls = []

    class Client:
        def __init__(self, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            pass

        async def post(self, *args, **kwargs):
            calls.append(1)
            if failure == "timeout":
                raise httpx.ReadTimeout("timeout")
            return httpx.Response(429 if failure == "rate_limit" else 401)

    async def no_sleep(*args):
        pass

    monkeypatch.setattr(httpx, "AsyncClient", Client)
    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    model = CompatibleProvider()
    with pytest.raises(AppError) as exc:
        await model.request("http://invalid", "test", "/test", {})
    assert len(calls) == (1 if failure == "unauthorized" else 3)
    assert exc.value.code in {"MODEL_TIMEOUT", "MODEL_UNAVAILABLE", "MODEL_CONFIGURATION"}
