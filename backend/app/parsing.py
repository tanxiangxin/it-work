import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader

from app.config import settings
from app.errors import AppError


@dataclass
class Section:
    content: str
    heading: str = ""
    page: int | None = None
    paragraph: int | None = None
    source_spans: list = field(default_factory=list)


def parse(path: Path, filename: str):
    ext = Path(filename).suffix.lower()
    sections = []
    try:
        if ext in {".md", ".txt"}:
            text = path.read_text(encoding="utf-8-sig")
            heading = []
            body = []
            start = 1
            fence = None
            for index, line in enumerate(text.splitlines(), 1):
                marker = re.match(r"^\s{0,3}(`{3,}|~{3,})", line) if ext == ".md" else None
                if marker:
                    if fence is None:
                        fence = marker[1]
                    elif marker[1][0] == fence[0] and len(marker[1]) >= len(fence):
                        fence = None
                    body.append(line)
                    continue
                match = re.match(r"^(#{1,6})\s+(.+)$", line) if ext == ".md" and fence is None else None
                if match:
                    if body:
                        sections.append(
                            Section("\n".join(body).strip(), " / ".join(heading), paragraph=start)
                        )
                    heading = heading[: len(match[1]) - 1] + [match[2]]
                    body, start = [], index + 1
                else:
                    if not (ext == ".md" and fence is None and line.startswith("> AI 联网整理资料")):
                        body.append(line)
            sections.append(Section("\n".join(body).strip(), " / ".join(heading), paragraph=start))
        elif ext == ".pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise AppError("ENCRYPTED_PDF", "不支持加密 PDF")
            if len(reader.pages) > settings().max_pdf_pages:
                raise AppError("PDF_TOO_LONG", "PDF 超过 100 页限制")
            sections = [
                Section(page.extract_text() or "", page=index) for index, page in enumerate(reader.pages, 1)
            ]
            if not any(s.content.strip() for s in sections):
                raise AppError("OCR_REQUIRED", "PDF 无可提取文本；首版不支持扫描件 OCR")
        elif ext == ".docx":
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise AppError("DOCX_TOO_LARGE", "DOCX 解压内容过大")
            word = WordDocument(path)
            heading = []
            for index, p in enumerate(word.paragraphs, 1):
                if p.style.name.startswith("Heading"):
                    digits = re.search(r"\d+", p.style.name)
                    level = int(digits[0]) if digits else 1
                    heading = heading[: level - 1] + [p.text]
                elif p.text.strip():
                    sections.append(Section(p.text, " / ".join(heading), paragraph=index))
            for index, table in enumerate(word.tables, 1):
                sections.append(
                    Section(
                        "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows),
                        heading=f"表格 {index}",
                    )
                )
        else:
            raise AppError("UNSUPPORTED_FILE", "只支持 MD、TXT、DOCX、PDF")
    except AppError:
        raise
    except Exception:
        raise AppError("PARSE_FAILED", "文件损坏、格式不匹配或编码不是 UTF-8") from None
    sections = [s for s in sections if s.content.strip() and s.heading.split(" / ")[-1] != "网页来源"]
    if not sections:
        raise AppError("EMPTY_DOCUMENT", "文件没有有效文本")
    if sum(len(s.content) for s in sections) > 2_000_000:
        raise AppError("TEXT_TOO_LARGE", "抽取文本超过 200 万字符限制")
    return sections


def protected_ranges(text):
    ranges, fence, fence_start, table_start, offset = [], None, None, None, 0
    for line in text.splitlines(keepends=True):
        marker = re.match(r"^\s{0,3}(`{3,}|~{3,})", line)
        if marker:
            if fence is None:
                fence, fence_start = marker[1], offset
            elif marker[1][0] == fence[0] and len(marker[1]) >= len(fence):
                ranges.append((fence_start, offset + len(line)))
                fence = None
        if fence is None and not marker:
            if re.match(r"^\s*\|.*\|\s*$", line):
                if table_start is None:
                    table_start = offset
            elif table_start is not None:
                ranges.append((table_start, offset))
                table_start = None
        offset += len(line)
    if fence_start is not None and fence is not None:
        ranges.append((fence_start, len(text)))
    if table_start is not None:
        ranges.append((table_start, len(text)))
    return ranges


def chunk_sections(sections):
    cfg = settings()
    packed = []
    for s in sections:
        text = s.content.strip()
        s = Section(
            text,
            s.heading,
            s.page,
            s.paragraph,
            [{"heading": s.heading, "page": s.page, "paragraph": s.paragraph, "start": 0, "end": len(text)}],
        )
        previous = packed[-1] if packed else None
        parent = s.heading.rpartition(" / ")[0]
        related = (
            previous
            and previous.page == s.page
            and (
                previous.heading == s.heading
                or (
                    parent and (previous.heading == parent or previous.heading.rpartition(" / ")[0] == parent)
                )
            )
        )
        protected = "```" in text or "~~~" in text or bool(re.search(r"^\s*\|.*\|\s*$", text, re.M))
        previous_protected = previous and (
            "```" in previous.content or "~~~" in previous.content or "|" in previous.content
        )
        separator = "\n\n" + (
            f"### {s.heading.split(' / ')[-1]}\n\n" if previous and previous.heading != s.heading else ""
        )
        if (
            related
            and not protected
            and not previous_protected
            and min(len(text), len(previous.content)) < cfg.chunk_min_chars
            and len(previous.content) + len(separator) + len(text) <= cfg.chunk_size
        ):
            previous.content += separator + text
            previous.source_spans.extend(s.source_spans)
        else:
            packed.append(s)
    chunks = []
    for s in packed:
        text = s.content
        protected = protected_ranges(text)
        start = 0
        while start < len(text):
            end = min(start + cfg.chunk_size, len(text))
            if end < len(text):
                # Prefer paragraph/sentence boundaries; never cut inside a fitting fenced block/table.
                boundaries = [
                    m.end() for m in re.finditer(r"\n\s*\n|[。！？]|(?<=[.!?])\s+", text[start:end])
                ]
                boundary = next(
                    (start + n for n in reversed(boundaries) if n >= min(800, cfg.chunk_size // 2)), -1
                )
                if boundary < 0:
                    boundary = text.rfind("\n", start + cfg.chunk_size // 2, end)
                for block_start, block_end in protected:
                    if block_start < end < block_end and block_end - block_start <= cfg.chunk_size:
                        if block_start > start:
                            boundary = block_start
                        elif block_end <= start + cfg.chunk_size:
                            boundary = block_end
                if boundary > start:
                    end = boundary
            spans = (
                s.source_spans
                if len(s.source_spans) > 1
                else [{**s.source_spans[0], "start": start, "end": end}]
            )
            chunks.append(Section(text[start:end], s.heading, s.page, s.paragraph, spans))
            if end == len(text):
                break
            # Natural small sections need no overlap; long-section overlap remains bounded.
            overlap = min(cfg.chunk_overlap, max(0, end - start - 300))
            if any(
                block_start == end and block_end - block_start + overlap > cfg.chunk_size
                for block_start, block_end in protected
            ):
                overlap = 0
            start = max(start + 1, end - overlap)
    return chunks
